import pytest

from core.parser import DocumentParser
from exceptions import DocumentParseError
from tests.conftest import FIXTURES


def test_parse_markdown_extracts_text_and_sections():
    parser = DocumentParser()
    document = parser.parse((FIXTURES / "sample_v1.md").read_bytes(), "sample_v1.md")

    assert document.file_type == "md"
    assert "SGD 옵티마이저" in document.content
    headings = [s.heading for s in document.sections]
    assert any("학습 알고리즘" in h for h in headings)


def test_parse_txt_handles_cp949():
    parser = DocumentParser()
    document = parser.parse("한글 인코딩 테스트".encode("cp949"), "note.txt")
    assert "한글 인코딩" in document.content


def test_split_sections_without_headings():
    parser = DocumentParser()
    sections = parser.split_sections("첫 문단입니다.\n\n둘째 문단입니다.")
    assert len(sections) == 2


def _docx_bytes() -> bytes:
    import io

    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_heading("설계 리뷰", level=1)
    document.add_paragraph("중앙 UPF 경유로 지연이 증가한다.")
    document.add_heading("결정 사항", level=2)
    document.add_paragraph("서비스 ID와 단말 위치를 함께 입력으로 사용한다.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "항목"
    table.cell(0, 1).text = "값"
    table.cell(1, 0).text = "평균 RTT"
    table.cell(1, 1).text = "15.9ms"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_parse_docx_extracts_paragraphs_headings_and_tables():
    parser = DocumentParser()
    document = parser.parse(_docx_bytes(), "design_review.docx")

    assert document.file_type == "docx"
    assert "중앙 UPF 경유로 지연이 증가한다." in document.content
    assert "# 설계 리뷰" in document.content  # Heading 1 → 마크다운 헤딩
    assert "## 결정 사항" in document.content
    assert "| 평균 RTT | 15.9ms |" in document.content  # 표도 읽는다

    headings = [s.heading for s in document.sections]
    assert "# 설계 리뷰" in headings and "## 결정 사항" in headings


def test_docx_is_listed_as_supported():
    from config import settings

    assert ".docx" in settings.SUPPORTED_EXT


def test_broken_docx_raises_e1005():
    parser = DocumentParser()
    with pytest.raises(DocumentParseError) as exc:
        parser.parse(b"not a real docx", "broken.docx")
    assert exc.value.error_code == "E1005"


def test_unsupported_extension_raises_e1001():
    parser = DocumentParser()
    with pytest.raises(DocumentParseError) as exc:
        parser.parse(b"data", "report.hwp")
    assert exc.value.error_code == "E1001"


def test_empty_content_raises():
    parser = DocumentParser()
    with pytest.raises(DocumentParseError):
        parser.parse(b"   \n  ", "empty.md")
