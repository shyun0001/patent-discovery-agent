"""ReportBuilder — 발명신고서 초안 렌더링 (LLM 사용 없음)."""
from __future__ import annotations

import io
from datetime import datetime

from config import settings
from core.elaborator import DisclosureDetail
from db.models import Document, DisclosureReport, InventionStructure, SearchQueries
from exceptions import ReportError
from patent.schemas import PriorArtSummary

_NO_PRIOR_ART = "선행기술 조사 미실시"


class ReportBuilder:
    def __init__(self, template_path=None):
        self.template_path = template_path or (settings.TEMPLATE_DIR / "disclosure_template.md")

    def build_disclosure(
        self,
        invention: InventionStructure,
        queries: SearchQueries | None = None,
        prior_art: PriorArtSummary | None = None,
        inventor_name: str = "",
        inventor_affiliation: str = "",
        detail: DisclosureDetail | None = None,
        documents: list[Document] | None = None,
    ) -> DisclosureReport:
        if not invention or not invention.problem:
            raise ReportError("발명 구조 정보가 비어 있습니다.", "E5002")

        queries = queries or SearchQueries(invention_id=invention.invention_id)
        detail = detail or DisclosureDetail()
        sections = {
            "document_no": f"IDF-{invention.invention_id[:8].upper()}",
            "title": invention.title or "(제목 없음)",
            "technical_field": invention.technical_field,
            "background_art": invention.background_art,
            "problem": invention.problem,
            "solution_means": invention.solution_means,
            "effect": invention.effect,
            "summary": invention.summary,
            "query_kr": queries.primary_kr,
            "query_en": queries.primary_en,
            "ipc_codes": ", ".join(
                f"{c.code}({c.description})" if c.description else c.code for c in queries.ipc_codes
            )
            or "-",
            "inventor_name": inventor_name,
            "inventor_affiliation": inventor_affiliation,
            "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        }
        sections.update(self._prior_art_fields(prior_art))
        sections.update(self._patentability_fields(invention))
        sections.update(self._detail_fields(detail, invention))
        sections["evidence"] = self._evidence(documents, prior_art)

        template = self._load_template()
        try:
            markdown = template.format(**sections)
        except KeyError as exc:
            raise ReportError(f"템플릿 필드가 누락되었습니다: {exc}", "E5002") from exc

        return DisclosureReport(
            invention_id=invention.invention_id,
            report_markdown=markdown,
            sections=sections,
            prior_art_included=bool(prior_art and prior_art.patents),
        )

    def build_from_draft(
        self,
        invention: InventionStructure,
        draft_markdown: str,
        queries: SearchQueries | None = None,
        prior_art: PriorArtSummary | None = None,
        documents: list[Document] | None = None,
    ) -> DisclosureReport:
        """LLM이 작성한 본문에 시스템이 보증하는 머리말과 근거 부록만 덧붙인다.

        본문은 손대지 않는다. 부록의 링크·커밋 SHA는 실제 수집 기록에서 생성하므로
        LLM이 지어낼 수 없는 값이다.
        """
        if not draft_markdown.strip():
            raise ReportError("발명신고서 본문이 비어 있습니다.", "E5002")

        header = self._header(invention, queries, prior_art)
        appendix = self._appendix(documents, prior_art)
        markdown = f"{header}\n\n---\n\n{draft_markdown.strip()}\n\n---\n\n{appendix}"

        return DisclosureReport(
            invention_id=invention.invention_id,
            report_markdown=markdown,
            sections={"body": draft_markdown.strip(), "header": header, "appendix": appendix},
            prior_art_included=bool(prior_art and prior_art.patents),
        )

    def _header(
        self,
        invention: InventionStructure,
        queries: SearchQueries | None,
        prior_art: PriorArtSummary | None,
    ) -> str:
        ipc = (
            ", ".join(code.code for code in queries.ipc_codes)
            if queries and queries.ipc_codes
            else "-"
        )
        risk = prior_art.overall_risk if prior_art and prior_art.patents else _NO_PRIOR_ART
        return "\n".join(
            [
                "# 발명신고서 (Invention Disclosure)",
                "",
                "| 항목 | 내용 |",
                "|------|------|",
                f"| 문서 번호 | `IDF-{invention.invention_id[:8].upper()}` |",
                f"| 작성일 | {datetime.now().strftime('%Y-%m-%d %H:%M')} |",
                f"| 기술 분야 | {invention.technical_field or '-'} |",
                f"| IPC 분류(안) | {ipc} |",
                f"| 선행기술 위험도 | {risk} |",
                "| 작성 | Patent Discovery Agent (AI 자동 생성 초안) |",
            ]
        )

    def _appendix(
        self, documents: list[Document] | None, prior_art: PriorArtSummary | None
    ) -> str:
        lines = ["## 부록 A. 분석 근거 (시스템 자동 기록)", "", self._evidence(documents, prior_art)]
        lines += [
            "",
            "## 부록 B. 문서 성격",
            "",
            "- 본 문서는 연구 산출물의 버전 변화를 분석해 자동 생성한 **초안**입니다.",
            "- 선행기술 조사 결과는 KIPRIS 키워드 검색과 초록 기반 유사도 평가이며, "
            "**정식 선행기술조사를 대체하지 않습니다.**",
            "- 출원 여부와 청구범위는 변리사 검토를 거쳐 결정하십시오.",
        ]
        return "\n".join(lines)

    def render_markdown(self, report: DisclosureReport) -> str:
        return report.report_markdown

    def render_prior_art_table(self, prior_art: PriorArtSummary, limit: int = 5) -> str:
        if not prior_art or not prior_art.patents:
            return "_검색된 선행기술이 없습니다._"

        header = (
            "| 출원번호 | 발명의 명칭 | 출원인 | 출원일 | 유사도 |\n"
            "|----------|-------------|--------|--------|--------|"
        )
        rows = []
        for patent in prior_art.patents[:limit]:
            assessment = prior_art.assessment_for(patent.application_number)
            score = f"{assessment.similarity_score:.2f}" if assessment else "-"
            title = self._cell((patent.invention_title or "")[:40])
            applicant = self._cell(patent.applicant_name)
            rows.append(
                f"| {patent.application_number} | {title} | {applicant} "
                f"| {patent.application_date} | {score} |"
            )
        return header + "\n" + "\n".join(rows)

    # ─── 확장 항목 렌더링 ──────────────────────────────
    @staticmethod
    def _patentability_fields(invention: InventionStructure) -> dict:
        score = invention.patentability
        return {
            "novelty": score.novelty or "-",
            "novelty_reasoning": score.novelty_reasoning or "-",
            "inventive_step": score.inventive_step or "-",
            "inventive_step_reasoning": score.inventive_step_reasoning or "-",
            "industrial_applicability": score.industrial_applicability or "-",
            "industrial_applicability_reasoning": score.industrial_applicability_reasoning or "-",
        }

    def _detail_fields(self, detail: DisclosureDetail, invention: InventionStructure) -> dict:
        if detail.components:
            rows = ["| # | 구성요소 | 기능 | 세부 동작 |", "|---|----------|------|-----------|"]
            for index, component in enumerate(detail.components, start=1):
                rows.append(
                    f"| {index} | {self._cell(component.name)} | {self._cell(component.function)} "
                    f"| {self._cell(component.detail)} |"
                )
            components_table = "\n".join(rows)
        else:
            components_table = "_구성요소 상세는 생성되지 않았습니다._"

        dependent = "\n\n".join(
            f"**청구항 {index}** {claim}"
            for index, claim in enumerate(detail.claims_dependent, start=2)
        ) or "_종속항 초안이 생성되지 않았습니다._"

        default_issues = "- 실험 데이터 보강 및 재현성 확인\n- 권리범위(청구항) 변리사 검토"
        return {
            "title_en": detail.title_en or "-",
            "purpose": detail.purpose or invention.problem,
            "components_table": components_table,
            "operation_steps": "\n".join(detail.operation)
            or "_동작 순서가 생성되지 않았습니다._",
            "embodiment": detail.embodiment or "_실시예가 생성되지 않았습니다._",
            "claim_independent": detail.claim_independent or "_청구항 초안이 생성되지 않았습니다._",
            "claims_dependent": dependent,
            "applications": "\n".join(f"- {item}" for item in detail.applications) or "-",
            "open_issues": "\n".join(f"- {item}" for item in detail.open_issues) or default_issues,
        }

    @staticmethod
    def _evidence(documents: list[Document] | None, prior_art: PriorArtSummary | None) -> str:
        """분석에 사용한 소스 리비전과 선행기술 링크를 남긴다."""
        lines = []
        for document in documents or []:
            source = document.source
            if not source:
                continue
            label = "이전 버전" if document.version == 1 else "변경 버전"
            ref = (source.source_ref or "")[:12]
            entry = f"- **{label}**: `{source.source_path}` @ `{ref}`"
            if source.source_url:
                entry += f" ([소스 보기]({source.source_url}))"
            lines.append(entry)

        if prior_art and prior_art.patents:
            lines.append("- **선행기술 원문**:")
            for patent in prior_art.patents[:3]:
                lines.append(
                    f"  - [{patent.application_number}]({patent.kipris_url}) "
                    f"{patent.invention_title[:40]}"
                )
        return "\n".join(lines) or "_근거 자료가 기록되지 않았습니다._"

    @staticmethod
    def _cell(text: str) -> str:
        """마크다운 표 셀 값 정리. KIPRIS는 복수 출원인을 '|'로 구분해 주므로 셀이 깨진다."""
        return (text or "").replace("|", ", ").replace("\n", " ").strip()

    def export_docx(self, report: DisclosureReport) -> bytes:
        """마크다운 초안을 간단한 DOCX로 변환 (Nice-to-Have)."""
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise ReportError(f"python-docx를 불러올 수 없습니다: {exc}", "E5001") from exc

        document = DocxDocument()
        for line in report.report_markdown.splitlines():
            stripped = line.strip()
            if stripped.startswith("### "):
                document.add_heading(stripped[4:], level=3)
            elif stripped.startswith("## "):
                document.add_heading(stripped[3:], level=2)
            elif stripped.startswith("# "):
                document.add_heading(stripped[2:], level=1)
            elif stripped:
                document.add_paragraph(stripped)

        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    # ─── 내부 ──────────────────────────────────────────
    def _load_template(self) -> str:
        try:
            return self.template_path.read_text(encoding="utf-8")
        except OSError as exc:
            raise ReportError(f"신고서 템플릿을 불러올 수 없습니다: {exc}", "E5001") from exc

    def _prior_art_fields(self, prior_art: PriorArtSummary | None) -> dict:
        if not prior_art or not prior_art.patents:
            return {
                "searched_db": "KIPRIS (국내 특허/실용신안)",
                "used_query": _NO_PRIOR_ART,
                "total_found": 0,
                "assessed_count": 0,
                "overall_risk": _NO_PRIOR_ART,
                "max_similarity": "-",
                "prior_art_table": "_선행기술 조사를 실시하지 않았거나 검색 결과가 없습니다._",
                "differentiating_points": "-",
            }

        differentiators = prior_art.key_differentiators or [
            a.differentiating_points for a in prior_art.assessments[:3] if a.differentiating_points
        ]
        note = " (※ 샘플 데이터)" if prior_art.is_sample else ""
        return {
            "searched_db": f"KIPRIS (국내 특허/실용신안){note}",
            "used_query": ", ".join(prior_art.used_queries) or "-",
            "total_found": prior_art.total_found,
            "assessed_count": prior_art.assessed_count,
            "overall_risk": prior_art.overall_risk,
            "max_similarity": f"{prior_art.max_similarity:.2f}",
            "prior_art_table": self.render_prior_art_table(prior_art),
            "differentiating_points": "\n".join(f"- {d}" for d in differentiators) or "-",
        }
