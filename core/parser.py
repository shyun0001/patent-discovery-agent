"""DocumentParser — 소스에서 가져온 파일 바이트에서 텍스트를 추출한다."""
from __future__ import annotations

import io
import re

from db.models import Document, Section, SourceRef
from exceptions import DocumentParseError
from utils.file_utils import ext_of

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")


class DocumentParser:
    # ─── 진입점 ────────────────────────────────────────
    def parse(
        self,
        file_bytes: bytes,
        filename: str,
        source: SourceRef | None = None,
        project_id: str = "",
        version: int = 1,
    ) -> Document:
        ext = ext_of(filename)
        extractors = {
            ".md": self.extract_text_from_md,
            ".txt": self.extract_text_from_txt,
            ".pdf": self.extract_text_from_pdf,
            ".pptx": self.extract_text_from_pptx,
            ".docx": self.extract_text_from_docx,
        }
        if ext not in extractors:
            raise DocumentParseError(
                "지원되는 형식: .md, .txt, .pdf, .pptx, .docx", "E1001"
            )

        text = self._clean_text(extractors[ext](file_bytes))
        if not text.strip():
            raise DocumentParseError(
                "문서에서 텍스트를 추출하지 못했습니다.",
                "E1003" if ext == ".pdf" else "E2001",
            )

        return Document(
            project_id=project_id,
            filename=filename,
            file_type=ext.lstrip("."),
            content=text,
            version=version,
            source=source,
            sections=self.split_sections(text),
        )

    # ─── 포맷별 추출 ───────────────────────────────────
    def extract_text_from_md(self, data: bytes) -> str:
        return self._decode(data)

    def extract_text_from_txt(self, data: bytes) -> str:
        return self._decode(data)

    def extract_text_from_pdf(self, data: bytes) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError(f"pypdf를 불러올 수 없습니다: {exc}", "E1003") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            pages = [(page.extract_text() or "") for page in reader.pages]
        except Exception as exc:
            raise DocumentParseError("PDF를 읽을 수 없습니다.", "E1003") from exc

        text = "\n\n".join(p.strip() for p in pages if p.strip())
        if not text.strip():
            raise DocumentParseError(
                "텍스트 추출이 불가능한 PDF입니다 (이미지 PDF일 수 있습니다).", "E1003"
            )
        return text

    def extract_text_from_pptx(self, data: bytes) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError(f"python-pptx를 불러올 수 없습니다: {exc}", "E1004") from exc
        try:
            presentation = Presentation(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError("PPTX를 읽을 수 없습니다.", "E1004") from exc

        chunks = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        line = "".join(run.text for run in paragraph.runs).strip()
                        if line:
                            texts.append(line)
            if texts:
                chunks.append(f"--- Slide {index} ---\n" + "\n".join(texts))
        return "\n\n".join(chunks)

    def extract_text_from_docx(self, data: bytes) -> str:
        """문단 + 표를 읽는다. Word 제목 스타일은 마크다운 헤딩으로 바꿔 섹션 분리에 태운다."""
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:  # pragma: no cover
            raise DocumentParseError(f"python-docx를 불러올 수 없습니다: {exc}", "E1005") from exc
        try:
            document = DocxDocument(io.BytesIO(data))
        except Exception as exc:
            raise DocumentParseError("DOCX를 읽을 수 없습니다.", "E1005") from exc

        chunks = []
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            level = self._heading_level(paragraph.style.name if paragraph.style else "")
            chunks.append(f"{'#' * level} {text}" if level else text)

        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                if any(cells):
                    rows.append("| " + " | ".join(cells) + " |")
            if rows:
                chunks.append("\n".join(rows))

        return "\n\n".join(chunks)

    @staticmethod
    def _heading_level(style_name: str) -> int:
        """'Heading 2' / '제목 2' → 2, 'Title' → 1, 그 외 0."""
        name = (style_name or "").strip()
        if name in ("Title", "제목"):
            return 1
        match = re.match(r"^(?:Heading|제목)\s*(\d)$", name)
        return min(int(match.group(1)), 6) if match else 0

    # ─── 섹션 분리 ─────────────────────────────────────
    def split_sections(self, text: str) -> list[Section]:
        """마크다운 헤딩 기준으로 분리하고, 헤딩이 없으면 빈 줄 기준으로 나눈다."""
        lines = text.splitlines()
        sections: list[Section] = []
        heading = ""
        buffer: list[str] = []

        def flush():
            if heading or any(line.strip() for line in buffer):
                sections.append(
                    Section(index=len(sections), heading=heading, content="\n".join(buffer).strip())
                )

        has_heading = any(_HEADING.match(line) for line in lines)
        if has_heading:
            for line in lines:
                match = _HEADING.match(line)
                if match:
                    flush()
                    heading = line.strip()
                    buffer = []
                else:
                    buffer.append(line)
            flush()
            return sections

        for block in re.split(r"\n\s*\n", text):
            block = block.strip()
            if block:
                title = block.splitlines()[0][:40]
                sections.append(Section(index=len(sections), heading=title, content=block))
        return sections

    # ─── 내부 ──────────────────────────────────────────
    @staticmethod
    def _decode(data: bytes) -> str:
        for encoding in ("utf-8", "utf-8-sig", "cp949", "euc-kr"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise DocumentParseError("파일 인코딩을 인식할 수 없습니다.", "E1002")

    @staticmethod
    def _clean_text(text: str) -> str:
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        text = re.sub(r"[ \t]+\n", "\n", text)
        return re.sub(r"\n{4,}", "\n\n\n", text).strip()
